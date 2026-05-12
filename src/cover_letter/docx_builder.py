"""DOCX rendering for generated cover letters."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO


def build_cover_letter_docx(
    *,
    candidate_name: str,
    candidate_location: str,
    candidate_email: str,
    candidate_phone: str,
    company_name: str,
    subject: str,
    greeting: str,
    paragraphs: list[str],
    closing: str,
    signature_name: str,
) -> bytes:
    """Build .docx bytes for the cover letter."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("DOCX generation requires python-docx. Please install dependencies.") from exc

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    if candidate_name.strip():
        p = document.add_paragraph()
        run = p.add_run(candidate_name.strip())
        run.bold = True
        run.font.size = Pt(12)

    if candidate_location.strip():
        document.add_paragraph(candidate_location.strip())

    contact_parts: list[str] = []
    if candidate_email.strip():
        contact_parts.append(candidate_email.strip())
    if candidate_phone.strip():
        contact_parts.append(candidate_phone.strip())
    if contact_parts:
        document.add_paragraph(" | ".join(contact_parts))

    document.add_paragraph(datetime.now().strftime("%B %d, %Y"))
    document.add_paragraph("Hiring Manager")
    document.add_paragraph(company_name.strip())

    subject_para = document.add_paragraph()
    subject_run = subject_para.add_run(subject.strip())
    subject_run.bold = True

    document.add_paragraph(greeting.strip())

    for item in paragraphs:
        para_text = str(item or "").strip()
        if not para_text:
            continue
        para = document.add_paragraph(para_text)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.add_paragraph(closing.strip())
    document.add_paragraph(signature_name.strip())

    output = BytesIO()
    document.save(output)
    return output.getvalue()

